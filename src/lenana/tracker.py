"""Read models, target history and exports shared by the dashboard and Telegram."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from html import escape
from io import BytesIO
from typing import Any

from sqlmodel import Session, col, select

from .config import Config
from .enums import OutboxKind, VerifyStatus
from .models import (
    Call,
    CallScore,
    Channel,
    CoinMention,
    ExtractionRun,
    MentionScan,
    Outbox,
    PriceCache,
    Transcript,
    Video,
    utcnow,
)
from .pipeline.mentions import normalise_coin
from .prices.coingecko import CoinGecko, resolve_symbol
from .utils import floor_hour, human_date, human_datetime


@dataclass
class TrackerData:
    leaderboard: list[dict[str, Any]] = field(default_factory=list)
    targets: list[dict[str, Any]] = field(default_factory=list)
    mentions: list[dict[str, Any]] = field(default_factory=list)
    timeline: list[dict[str, Any]] = field(default_factory=list)
    prices: list[dict[str, Any]] = field(default_factory=list)
    videos: int = 0
    scanned: int = 0
    transcripts: int = 0


def load_tracker(
    session: Session,
    cfg: Config,
    days: int | None = None,
    channels: list[int] | None = None,
    include_sponsored: bool = False,
    now: datetime | None = None,
) -> TrackerData:
    now = now or utcnow()
    cutoff = now - timedelta(days=days) if days else None
    video_rows = session.exec(select(Video)).all()
    videos = {
        v.id: v
        for v in video_rows
        if (cutoff is None or v.published_at >= cutoff)
        and (channels is None or v.channel_pk in channels)
    }
    names = {c.id: c.name or c.handle for c in session.exec(select(Channel)).all()}
    active = {r.id for r in session.exec(select(ExtractionRun)).all() if r.is_active}
    scores = {s.call_pk: s for s in session.exec(select(CallScore)).all()}
    prices = session.exec(select(PriceCache).order_by(col(PriceCache.ts_hour))).all()
    current = {p.coingecko_id: p for p in prices if p.source == "simple_price"}
    data = TrackerData(videos=len(videos))
    data.scanned = sum(s.video_pk in videos for s in session.exec(select(MentionScan)).all())
    data.transcripts = sum(t.video_pk in videos for t in session.exec(select(Transcript)).all())
    evidence: dict[tuple[str, int], dict[str, Any]] = {}

    def source(symbol: str, video: Video, at: int, quote: str, sponsored: bool) -> dict[str, Any]:
        spoken = video.published_at + timedelta(seconds=at)
        return {
            "Coin": symbol,
            "Mention date": spoken.astimezone(cfg.settings.tz).isoformat(),
            "Channel": names.get(video.channel_pk, "Unknown"),
            "Video": video.title,
            "Source": f"{video.url}&t={at}s",
            "Quote": quote,
            "Sponsored": sponsored,
            "_video_pk": video.id,
            "_published": video.published_at,
            "_channel_pk": video.channel_pk,
        }

    for mention in session.exec(select(CoinMention)).all():
        video = videos.get(mention.video_pk)
        if video is None or (mention.is_sponsored and not include_sponsored):
            continue
        symbol = normalise_coin(mention.token_symbol)
        evidence[(symbol, mention.video_pk)] = source(
            symbol, video, mention.timestamp_seconds, mention.quote, mention.is_sponsored
        )

    for call in session.exec(select(Call)).all():
        video = videos.get(call.video_pk)
        if (
            video is None
            or call.extraction_run_pk not in active
            or (call.verify_status != VerifyStatus.VERIFIED)
            or (call.is_sponsored and not include_sponsored)
        ):
            continue
        symbol = normalise_coin(call.token_symbol)
        evidence.setdefault(
            (symbol, call.video_pk),
            source(symbol, video, call.timestamp_seconds, call.quote, call.is_sponsored),
        )
        spoken = call.entry_price_at or video.published_at + timedelta(
            seconds=call.timestamp_seconds
        )
        due = spoken + timedelta(days=call.timeframe_days) if call.timeframe_days else None
        score = scores.get(call.id) if call.id is not None else None
        if call.price_target is None:
            status = "No numeric target"
        elif score and score.target_touched is True:
            status = "Reached"
        elif score and score.scored_at:
            status = "Not reached" if score.target_touched is False else "Evaluated — touch unknown"
        elif due and due <= now:
            status = "Overdue — awaiting evaluation"
        elif due:
            status = "Pending"
        else:
            status = "No deadline stated"
        point = current.get(call.coingecko_id or "")
        usd_target = call.price_target if (call.target_currency or "").upper() == "USD" else None
        distance = (
            (usd_target - point.price_usd) / point.price_usd * 100
            if point and usd_target is not None
            else None
        )
        change = (
            (point.price_usd - call.entry_price_usd) / call.entry_price_usd * 100
            if point and call.entry_price_usd
            else None
        )
        row: dict[str, Any] = {
            "Coin": symbol,
            "Prediction date": spoken.astimezone(cfg.settings.tz).isoformat(),
            "Target": call.price_target,
            "Currency": call.target_currency or "Unspecified",
            "Direction": str(call.direction),
            "Deadline": due.astimezone(cfg.settings.tz).isoformat() if due else None,
            "Timeframe stated": call.timeframe_text,
            "Deadline estimated": due is not None,
            "Status": status,
            "Current USD": point.price_usd if point else None,
            "Price at call USD": call.entry_price_usd,
            "Change since call %": change,
            "Distance to target %": distance,
            "Price updated": point.fetched_at.astimezone(cfg.settings.tz).isoformat()
            if point
            else None,
            "Channel": names.get(video.channel_pk, "Unknown"),
            "Speaker": call.speaker,
            "Conditional": call.is_conditional,
            "Sponsored": call.is_sponsored,
            "Quote": call.quote,
            "Source": f"{video.url}&t={call.timestamp_seconds}s",
            "_call_pk": call.id,
            "_coin_id": call.coingecko_id,
            "_created": call.created_at,
            "_scored": score.scored_at if score else None,
        }
        data.targets.append(row)
    data.targets.sort(key=lambda r: r["Prediction date"], reverse=True)
    data.mentions = sorted(evidence.values(), key=lambda r: r["Mention date"], reverse=True)
    by_coin: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for mention_row in data.mentions:
        by_coin[mention_row["Coin"]].append(mention_row)
    for symbol, rows in by_coin.items():
        recent = sum(r["_published"] >= now - timedelta(days=7) for r in rows)
        previous = sum(
            now - timedelta(days=14) <= r["_published"] < now - timedelta(days=7) for r in rows
        )
        targets = [t for t in data.targets if t["Coin"] == symbol]
        coin_id, mapping = resolve_symbol(session, symbol)
        point = current.get(coin_id or "") if mapping == "mapped" else None
        data.leaderboard.append(
            {
                "Coin": symbol,
                "Videos mentioning": len(rows),
                "Channels": len({r["_channel_pk"] for r in rows}),
                "With numeric targets": sum(t["Target"] is not None for t in targets),
                "Last 7 days": recent,
                "Previous 7 days": previous,
                "Weekly change": recent - previous,
                "Latest mention": max(r["Mention date"] for r in rows),
                "Current USD": point.price_usd if point else None,
                "Price updated": point.fetched_at.astimezone(cfg.settings.tz).isoformat()
                if point
                else None,
            }
        )
    data.leaderboard.sort(key=lambda r: (-r["Videos mentioning"], -r["Channels"], r["Coin"]))
    dates = sorted({t["Prediction date"][:10] for t in data.targets})
    for coin in data.leaderboard:
        row = {"Coin": coin["Coin"], "Current USD": coin["Current USD"]}
        has_target = False
        for day in dates:
            matches = [
                t
                for t in data.targets
                if t["Coin"] == coin["Coin"]
                and t["Prediction date"].startswith(day)
                and t["Target"] is not None
            ]
            values = [
                f"{t['Target']:g} {t['Currency']} · {t['Channel']} "
                f"· {human_datetime(t['Prediction date'])} · "
                f"{t['Timeframe stated'] or 'no deadline'}"
                for t in matches
            ]
            row[day] = " | ".join(values) or "—"
            has_target = has_target or bool(values)
        row["Targets"] = "See dated targets" if has_target else "No target stated"
        data.timeline.append(row)
    coin_symbols = {t["_coin_id"]: t["Coin"] for t in data.targets if t["_coin_id"]}
    for symbol in by_coin:
        coin_id, mapping = resolve_symbol(session, symbol)
        if coin_id and mapping == "mapped":
            coin_symbols.setdefault(coin_id, symbol)
    data.prices = [
        {
            "Coin": coin_symbols[p.coingecko_id],
            "Date": p.ts_hour.astimezone(cfg.settings.tz).isoformat(),
            "Price USD": p.price_usd,
            "Source": p.source,
        }
        for p in prices
        if p.coingecko_id in coin_symbols
    ]
    return data


def public_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    date_fields = {
        "Mention date",
        "Prediction date",
        "Deadline",
        "Price updated",
        "Latest mention",
        "Date",
    }
    result = []
    for row in rows:
        visible = {}
        for key, value in row.items():
            if key.startswith("_"):
                continue
            label = (
                human_date(key) if len(key) == 10 and key[4:5] == "-" and key[7:8] == "-" else key
            )
            visible[label] = human_datetime(value) if key in date_fields and value else value
        result.append(visible)
    return result


def workbook_bytes(data: TrackerData, timezone: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    book.remove(book.active)
    for name, rows in (
        ("Coin leaderboard", data.leaderboard),
        ("Target timeline", data.timeline),
        ("Predictions", data.targets),
        ("Coin mentions", data.mentions),
        ("Price history", data.prices),
    ):
        sheet = book.create_sheet(name)
        rows = public_rows(rows)
        if not rows:
            sheet.append(["No data for this selection"])
            continue
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(h) for h in headers])
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173C36")
        # Quotes, channel names and symbols are external input, never formulas.
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if isinstance(cell.value, str) and cell.value.startswith(
                    "https://www.youtube.com/"
                ):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet.columns:
            sheet.column_dimensions[cells[0].column_letter].width = min(
                max(len(str(c.value or "")) for c in cells) + 2, 65
            )
    info = book.create_sheet("Read me")
    for line in (
        f"Dates use {timezone}.",
        "Mention counts are distinct videos; sponsored segments are excluded by default.",
        "Daily columns are the dates predictions were spoken, not their deadlines.",
        "Each prediction is preserved separately. Targets are never averaged.",
        "Deadlines are estimates; see the original timeframe and quote.",
        "Reached status requires a recorded evaluation, not merely a current quote.",
        f"Mention scan coverage: {data.scanned}/{data.transcripts} saved transcripts.",
        "A blank current price means the coin was not mapped or a quote is unavailable.",
    ):
        info.append([line])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def refresh_current_prices(session: Session, cfg: Config) -> int:
    symbols = {normalise_coin(m.token_symbol) for m in session.exec(select(CoinMention)).all()}
    symbols.update(
        normalise_coin(c.token_symbol)
        for c in session.exec(select(Call).where(Call.verify_status == VerifyStatus.VERIFIED)).all()
    )
    ids = {
        coin_id
        for symbol in symbols
        if (resolution := resolve_symbol(session, symbol))[1] == "mapped"
        and (coin_id := resolution[0]) is not None
    }
    if not ids:
        return 0
    api = CoinGecko(cfg)
    updated = 0
    now = utcnow()
    for start in range(0, len(ids), 100):
        payload = api._get(
            "/simple/price",
            {"ids": ",".join(sorted(ids)[start : start + 100]), "vs_currencies": "usd"},
        )
        if not isinstance(payload, dict):
            continue
        for coin_id, quote in payload.items():
            price = quote.get("usd") if isinstance(quote, dict) else None
            if not isinstance(price, int | float) or price <= 0:
                continue
            point = session.exec(
                select(PriceCache).where(
                    PriceCache.coingecko_id == coin_id, PriceCache.ts_hour == floor_hour(now)
                )
            ).first()
            if point is None:
                point = PriceCache(coingecko_id=coin_id, ts_hour=floor_hour(now), price_usd=price)
            point.price_usd = price
            point.source = "simple_price"
            point.fetched_at = now
            session.add(point)
            updated += 1
        session.commit()
    return updated


def queue_daily_report(session: Session, cfg: Config, force: bool = False) -> int:
    from .notify.outbox import enqueue

    now = utcnow()
    day = now.astimezone(cfg.settings.tz).date().isoformat()
    key = f"tracker-report:{day}" + (f":{now.isoformat()}" if force else "")
    if session.exec(select(Outbox).where(Outbox.dedupe_key == key)).first():
        return 0
    data = load_tracker(session, cfg)
    if not data.mentions:
        return 0
    recent = [
        t
        for t in data.targets
        if t["_created"].astimezone(cfg.settings.tz).date().isoformat() == day
    ]
    reached = [
        t
        for t in data.targets
        if t["Status"] == "Reached"
        and t["_scored"]
        and t["_scored"].astimezone(cfg.settings.tz).date().isoformat() == day
    ]
    ranked = sorted(data.leaderboard, key=lambda r: (-r["Last 7 days"], -r["Videos mentioning"]))
    top = [r for r in ranked if r["Last 7 days"]][:10]
    scope = "last 7 days"
    if not top:
        top, scope = data.leaderboard[:10], "all saved videos; no mentions in the last 7 days"
    lines = [f"📊 <b>Coin tracker · {human_date(day)}</b>", f"Most discussed ({scope}):"]
    for row in top:
        count = row["Last 7 days"] if scope == "last 7 days" else row["Videos mentioning"]
        lines.append(
            f"{escape(row['Coin'])}: {count} videos · "
            f"{row['Channels']} channels overall · weekly change {row['Weekly change']:+d}"
        )
    lines.extend(
        [
            f"\nNewly processed predictions: {len(recent)}",
            f"Targets confirmed reached today: {len(reached)}",
            f"Mention scan coverage: {data.scanned}/{data.transcripts} saved transcripts.",
            "Full target history and source quotes are in the attached spreadsheet.",
        ]
    )
    numeric_recent = [t for t in recent if t["Target"] is not None]
    if numeric_recent:
        lines.append("\n<b>New targets (up to 10; full list in spreadsheet)</b>")
        for target in numeric_recent[:10]:
            lines.append(
                f"{escape(target['Coin'])} → {target['Target']:g} "
                f"{escape(target['Currency'])} · {escape(target['Channel'])} · "
                f"{escape(target['Timeframe stated'] or 'no deadline stated')}"
            )
    if reached:
        lines.append("\n<b>Confirmed reached</b>")
        lines.extend(
            f"{escape(t['Coin'])}: {t['Target']:g} {escape(t['Currency'])}" for t in reached[:10]
        )
    path = cfg.settings.data_dir / "reports" / f"coin-tracker-{day}-{now:%H%M%S%f}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(workbook_bytes(data, cfg.settings.display_tz))
    rows = enqueue(session, key, OutboxKind.DAILY_REPORT, "\n".join(lines), cfg)
    rows += enqueue(
        session,
        key + ":xlsx",
        OutboxKind.DAILY_REPORT,
        f"Coin tracker · {human_date(day)} · full leaderboard, targets and source quotes",
        cfg,
    )
    rows[-1].document_path = str(path.resolve())
    session.commit()
    return len(rows)
