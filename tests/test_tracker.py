from __future__ import annotations

from datetime import timedelta
from io import BytesIO

import httpx
import pytest
import respx
from openpyxl import load_workbook
from sqlmodel import Session, select

from yt_crypto_signals.config import Config
from yt_crypto_signals.enums import (
    Direction,
    OutboxKind,
    OutboxState,
    TranscriptProvider,
    VerifyStatus,
)
from yt_crypto_signals.models import (
    Call,
    CallScore,
    Channel,
    CoinMention,
    ExtractionRun,
    Outbox,
    Transcript,
    Video,
    utcnow,
)
from yt_crypto_signals.notify.outbox import send_pending
from yt_crypto_signals.notify.telegram import TelegramClient
from yt_crypto_signals.pipeline.mentions import (
    MentionExtraction,
    RawMention,
    checked_mentions,
    mentions,
)
from yt_crypto_signals.tracker import load_tracker, queue_daily_report, workbook_bytes


def saved_video(session: Session, suffix: str = "1", days: int = 1) -> Video:
    channel = Channel(handle=f"@channel{suffix}", name=f"Channel {suffix}")
    session.add(channel)
    session.commit()
    assert channel.id is not None
    video = Video(
        video_id=f"video{suffix}",
        channel_pk=channel.id,
        published_at=utcnow() - timedelta(days=days),
    )
    session.add(video)
    session.commit()
    return video


def prediction(session: Session, video: Video, target: float, active: bool = True) -> Call:
    assert video.id is not None
    run = ExtractionRun(video_pk=video.id, prompt_version="v1", model="test", is_active=active)
    session.add(run)
    session.commit()
    assert run.id is not None
    call = Call(
        video_pk=video.id,
        extraction_run_pk=run.id,
        token_symbol="BTC",
        direction=Direction.UP,
        price_target=target,
        target_currency="USD",
        speaker="Host",
        verify_status=VerifyStatus.VERIFIED,
        timeframe_days=7,
        quote="Bitcoin will rise to the price target next week",
    )
    session.add(call)
    session.commit()
    return call


def test_mentions_without_predictions_and_duplicate_votes(session: Session, config: Config) -> None:
    video = saved_video(session)
    assert video.id is not None
    session.add(
        CoinMention(
            video_pk=video.id, token_symbol="BTC", quote="Bitcoin discussed", timestamp_seconds=1
        )
    )
    session.add(
        CoinMention(
            video_pk=video.id, token_symbol="ETH", quote="Ethereum discussed", timestamp_seconds=2
        )
    )
    session.add(
        CoinMention(
            video_pk=video.id,
            token_symbol="AD",
            quote="Sponsored coin",
            timestamp_seconds=3,
            is_sponsored=True,
        )
    )
    session.commit()
    prediction(session, video, 120000)
    prediction(session, video, 130000)
    prediction(session, video, 999999, active=False)
    data = load_tracker(session, config)
    rows = {r["Coin"]: r for r in data.leaderboard}
    assert set(rows) == {"BTC", "ETH"}
    assert rows["BTC"]["Videos mentioning"] == 1
    assert rows["BTC"]["With numeric targets"] == 2
    assert rows["ETH"]["With numeric targets"] == 0
    assert len(data.targets) == 2
    btc_grid = next(r for r in data.timeline if r["Coin"] == "BTC")
    assert "120000" in str(btc_grid) and "130000" in str(btc_grid)
    assert "999999" not in str(btc_grid)
    assert next(r for r in data.timeline if r["Coin"] == "ETH")["Targets"] == "No target stated"
    assert not load_tracker(session, config, channels=[]).leaderboard
    assert "AD" in {
        r["Coin"] for r in load_tracker(session, config, include_sponsored=True).leaderboard
    }


def test_trends_use_publish_time_and_status_requires_evaluation(
    session: Session, config: Config
) -> None:
    recent, previous = saved_video(session, "1", 1), saved_video(session, "2", 8)
    call = prediction(session, recent, 120000)
    prediction(session, previous, 110000)
    data = load_tracker(session, config)
    assert data.leaderboard[0]["Last 7 days"] == 1
    assert data.leaderboard[0]["Previous 7 days"] == 1
    assert data.leaderboard[0]["Weekly change"] == 0
    assert {t["Status"] for t in data.targets} == {"Pending", "Overdue — awaiting evaluation"}
    assert call.id is not None
    session.add(
        CallScore(call_pk=call.id, due_at=utcnow(), scored_at=utcnow(), target_touched=True)
    )
    session.commit()
    assert "Reached" in {t["Status"] for t in load_tracker(session, config).targets}
    assert load_tracker(session, config, days=7).leaderboard[0]["Videos mentioning"] == 1


def test_evidence_rejects_hallucinated_coin_and_merges_aliases(config: Config) -> None:
    text = "Bitcoin is an interesting project and we discuss its adoption today"
    transcript = Transcript(
        video_pk=1,
        provider=TranscriptProvider.YOUTUBE_TRANSCRIPT_API,
        segments=[{"start": 0, "duration": 10, "text": text}],
    )
    items = [
        RawMention(
            token_symbol=ticker,
            token_name="Bitcoin",
            quote=text,
            timestamp_seconds=0,
            is_sponsored=False,
        )
        for ticker in ("BTC", "XBT")
    ]
    items.append(
        RawMention(
            token_symbol="FAKE",
            token_name="Fakecoin",
            quote="Fakecoin is an interesting project and we discuss its adoption today",
            timestamp_seconds=0,
            is_sponsored=False,
        )
    )
    accepted, rejected = checked_mentions(items, transcript, config)
    assert [item.token_symbol for item in accepted] == ["BTC"]
    assert rejected == 1


def test_empty_scans_are_idempotent_and_changed_transcript_rescans(
    session: Session,
    config: Config,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from yt_crypto_signals.llm.client import Usage
    from yt_crypto_signals.pipeline import mentions as module

    video = saved_video(session)
    assert video.id is not None
    transcript = Transcript(
        video_pk=video.id,
        provider=TranscriptProvider.YOUTUBE_TRANSCRIPT_API,
        segments=[{"start": 0, "duration": 10, "text": "General market discussion today"}],
    )
    session.add(transcript)
    session.commit()

    class FakeExtractor:
        model = "test"

        def __init__(self, _cfg: Config) -> None:
            pass

        def extract(self, *args: object) -> tuple[MentionExtraction, Usage]:
            return MentionExtraction(mentions=[]), Usage()

    monkeypatch.setattr(module, "ClaudeExtractor", FakeExtractor)
    assert mentions(session, config).processed == 1
    assert mentions(session, config).processed == 0
    transcript.segments = [{"start": 0, "duration": 10, "text": "Updated market discussion today"}]
    session.add(transcript)
    session.commit()
    assert mentions(session, config).processed == 1


def test_workbook_and_daily_report_are_safe_and_deduplicated(
    session: Session, config: Config
) -> None:
    video = saved_video(session)
    assert video.id is not None
    session.add(
        CoinMention(
            video_pk=video.id, token_symbol="ETH", quote='=HYPERLINK("bad")', timestamp_seconds=0
        )
    )
    session.commit()
    data = load_tracker(session, config)
    book = load_workbook(BytesIO(workbook_bytes(data, "UTC")))
    assert "Target timeline" in book.sheetnames
    quote_column = next(c.column for c in book["Coin mentions"][1] if c.value == "Quote")
    assert book["Coin mentions"].cell(2, quote_column).data_type == "s"
    assert queue_daily_report(session, config) == 2
    assert queue_daily_report(session, config) == 0
    rows = session.exec(select(Outbox)).all()
    assert len(rows) == 2
    assert sum(row.document_path is not None for row in rows) == 1


@respx.mock
def test_document_delivery_uses_outbox_and_handles_non_json(
    session: Session,
    config: Config,
) -> None:
    path = config.settings.data_dir / "report.xlsx"
    path.write_bytes(b"spreadsheet")
    route = respx.post(
        f"https://api.telegram.org/bot{config.settings.telegram_bot_token}/sendDocument"
    )
    route.mock(return_value=httpx.Response(200, json={"ok": True, "result": {"message_id": 10}}))
    row = Outbox(
        dedupe_key="doc",
        kind=OutboxKind.DAILY_REPORT,
        chat_id=config.settings.telegram_chat_id,
        body="Tracker",
        document_path=str(path),
    )
    session.add(row)
    session.commit()
    assert send_pending(session, config, sleep=False).sent == 1
    assert row.state == OutboxState.SENT
    assert b"spreadsheet" in route.calls[0].request.content
    assert send_pending(session, config, sleep=False).sent == 0
    route.mock(return_value=httpx.Response(502, text="Bad Gateway"))
    assert TelegramClient(config).send_document(path, "test").is_retryable
